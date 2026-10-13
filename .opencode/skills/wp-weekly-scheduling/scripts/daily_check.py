# -*- coding: utf-8 -*-
"""WP DDL 每日弹窗检查（由 generate_schedule.py 自动生成）。"""
import argparse
import datetime as dt
import glob
import os
import subprocess

DEFAULT_OUTPUT_DIR = r"C:\AIProjects\3-wp-weekly-scheduling\.opencode\skills\wp-weekly-scheduling\output"


def find_latest(output_dir):
    files = [p for p in glob.glob(os.path.join(output_dir, "WP周排班_*.xlsx"))
             if not os.path.basename(p).startswith("~$")]
    return max(files, key=os.path.getmtime) if files else None


def load_reminders(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if "日历提醒汇总" not in wb.sheetnames:
        return []
    ws = wb["日历提醒汇总"]
    rows = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r[0] is None:
            continue
        d = r[0]
        if isinstance(d, dt.datetime):
            d = d.date()
        elif isinstance(d, str):
            try:
                d = dt.datetime.strptime(d[:10], "%Y-%m-%d").date()
            except ValueError:
                continue
        else:
            continue
        rows.append({"date": d, "title": r[3] or "", "detail": r[4] or ""})
    return rows


def notify(title, msg):
    try:
        from plyer import notification
        notification.notify(title=title, message=msg, timeout=15)
        return
    except Exception:
        pass
    try:
        from win10toast import ToastNotifier
        ToastNotifier().show_toast(title, msg, duration=15, threaded=False)
        return
    except Exception:
        pass
    ps = ("Add-Type -AssemblyName System.Windows.Forms; "
          "[System.Windows.Forms.MessageBox]::Show(@'\n" + msg + "\n'@, '" + title + "')")
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=False)
    except Exception:
        print(title)
        print(msg)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    ap.add_argument("--days", type=int, default=7)
    args = ap.parse_args()
    latest = find_latest(args.output_dir)
    if not latest:
        return
    today = dt.date.today()
    end = today + dt.timedelta(days=args.days)
    rows = [r for r in load_reminders(latest) if today <= r["date"] <= end]
    if not rows:
        return
    rows.sort(key=lambda r: r["date"])
    lines = []
    for r in rows:
        lines.append("{}  {}".format(r["date"].isoformat(), r["title"]))
        if r["detail"]:
            lines.append(r["detail"])
    notify("WP DDL 提醒", "\n".join(lines))


if __name__ == "__main__":
    main()
