import os

import openpyxl

import todo_service
from config import load_config
from db import get_conn, JOB_LOCK

NODE_SHEET = "项目关键节点"


def _to_str(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _col(header, name):
    for i, h in enumerate(header):
        if h == name or h.replace(" ", "") == name.replace(" ", ""):
            return i
    return None


def read_ddl_rows():
    cfg = load_config()
    path = cfg["datasources"].get("schedule_output")
    if not path or not os.path.exists(path):
        raise FileNotFoundError("排班输出文件不存在，请在“数据源设置”重新指定")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = None
        for name in wb.sheetnames:
            if NODE_SHEET in name:
                sheet = name
                break
        if not sheet:
            raise ValueError("排班输出里找不到“项目关键节点”sheet")
        ws = wb[sheet]
        header = None
        header_idx = 1
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            if row and any(v is not None and str(v).strip() for v in row):
                header = [str(v).strip() if v is not None else "" for v in row]
                header_idx = i
                break
        c_code = _col(header, "Engagement Code")
        c_name = _col(header, "Engagement Name")
        c_node = _col(header, "节点类型")
        c_date = _col(header, "日期")
        if None in (c_code, c_node, c_date):
            raise ValueError("“项目关键节点”缺少 Engagement Code / 节点类型 / 日期 列")
        rows = []
        for row in ws.iter_rows(min_row=header_idx + 1, values_only=True):
            if row is None or c_code >= len(row):
                continue
            code = row[c_code]
            if code is None or str(code).strip() == "":
                continue
            name = row[c_name] if c_name is not None and c_name < len(row) else ""
            node = row[c_node] if c_node < len(row) else ""
            d = row[c_date] if c_date < len(row) else ""
            rows.append({
                "code": str(code).strip(),
                "name": _to_str(name),
                "node": _to_str(node),
                "date": _to_str(d),
            })
        return rows
    finally:
        wb.close()


def refresh_ddl():
    with JOB_LOCK:
        return _refresh_ddl()


def _refresh_ddl():
    rows = read_ddl_rows()
    conn = get_conn()
    added = updated = 0
    try:
        for r in rows:
            if not r["node"] or not r["date"]:
                continue
            title = (r["name"] + " " + r["node"]).strip()
            note = f"项目：{r['name']}（Code{r['code']}）| 节点：{r['node']}"
            existing = conn.execute(
                "SELECT id, archived_at FROM todos WHERE source_code=? AND source_node=?",
                (r["code"], r["node"]),
            ).fetchone()
            if existing:
                if existing["archived_at"] is None:
                    conn.execute(
                        "UPDATE todos SET title=?, due_date=?, note=?, category='DDL' WHERE id=?",
                        (title, r["date"], note, existing["id"]),
                    )
                    updated += 1
            else:
                conn.execute(
                    "INSERT INTO todos(title, status, due_date, category, quadrant, "
                    "repeat_type, repeat_param, note, source_code, source_node) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (title, todo_service.STATUS_TODO, r["date"], "DDL",
                     todo_service.QUADRANT_1, "none", "", note,
                     r["code"], r["node"]),
                )
                added += 1
        conn.commit()
    finally:
        conn.close()
    todo_service.apply_status_rules()
    return {"added": added, "updated": updated, "total": len(rows)}
