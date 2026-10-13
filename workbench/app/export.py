import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import timesheet_service as ts
from config import EXPORT_DIR

HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FILL = PatternFill("solid", fgColor="E2EFDA")
OVER_FONT = Font(color="C00000", bold=True, name="Microsoft YaHei")
BASE_FONT = Font(name="Microsoft YaHei")
HEAD_FONT = Font(name="Microsoft YaHei", bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

_INVALID = re.compile(r'[\\/:*?"<>|\s]+')


def _safe(text):
    return _INVALID.sub("_", str(text or "").strip())


def _filename(mode, codes, managers):
    stamp = date.today().strftime("%Y%m%d")
    if mode == "person":
        names = [m for m in (managers or []) if m and m.strip()]
        first = _safe(names[0]) if names else "结果"
        suffix = "等" if len(names) > 1 else ""
    else:
        names = [c for c in (codes or []) if c and str(c).strip()]
        first = _safe(names[0]) if names else "结果"
        suffix = "等" if len(names) > 1 else ""
    return f"工时筛选_{stamp}_{first}{suffix}.xlsx"


def export_timesheet(mode, codes, managers):
    result = ts.build_result(mode, codes, managers, limit=None)
    if not result["rows"]:
        raise ValueError("没有可导出的数据，请重新筛选")
    wb = Workbook()
    ws = wb.active
    ws.title = "工时筛选"

    headers = ["项目名称", "code", "staffname", "manager name", "hours", "期间", "描述"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for r in result["rows"]:
        ws.append([
            r["project_name"], r["project_code"], r["staff"],
            r["manager"], r["hours"], r["period"], r["description"],
        ])

    ws.append([])
    ws.append(["项目合计"])
    ws.cell(row=ws.max_row, column=1).font = Font(name="Microsoft YaHei", bold=True, size=12)
    ws.append(["项目名称", "code", "合计工时", "预算 (Outlook Hours)", "是否超预算"])
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = HEAD_FONT
        cell.fill = TOTAL_FILL
        cell.border = BORDER
    for t in result["totals"]:
        flag = "—"
        if t["budget"] is not None:
            flag = "超预算" if t["over"] else "未超"
        ws.append([t["project_name"], t["project_code"], t["hours"], t["budget"], flag])
        if t["over"]:
            for c in range(1, 6):
                ws.cell(row=ws.max_row, column=c).font = OVER_FONT

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.font is None or cell.font.name is None:
                cell.font = BASE_FONT
            cell.border = BORDER
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "0.00"

    widths = [34, 12, 18, 20, 10, 14, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    filename = _filename(mode, codes, managers)
    path = EXPORT_DIR / filename
    wb.save(path)
    return {"path": str(path), "filename": filename, "rows": len(result["rows"])}
