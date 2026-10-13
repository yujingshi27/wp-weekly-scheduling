import os
import time
import zipfile
from datetime import datetime, timedelta

import openpyxl
from lxml import etree

from config import load_config
from db import get_conn, JOB_LOCK
from timesheet_service import norm_code

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

TIMESHEET_KEYS = [
    ("project_code", "项目编号列"),
    ("staff", "员工列"),
    ("manager", "manager 列"),
    ("hours", "工时列"),
    ("period", "期间列"),
    ("description", "描述列"),
    ("project_name", "项目名称列"),
]

_BUILTIN_DATE_FMT = set(list(range(14, 23)) + list(range(45, 48)) + [27, 30, 36, 50, 57])


# ------------------------- small-file helpers (openpyxl) -------------------------
def list_sheets(path):
    if not path or not os.path.exists(path):
        raise FileNotFoundError("文件不存在：" + str(path))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def find_sheet(sheet_names, keyword, exact=False):
    for name in sheet_names:
        if (name == keyword) if exact else (keyword in name):
            return name
    return None


def read_headers(path, sheet=None):
    if not path or not os.path.exists(path):
        raise FileNotFoundError("文件不存在：" + str(path))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            if row and any(v is not None and str(v).strip() for v in row):
                return [str(v).strip() if v is not None else "" for v in row]
        return []
    finally:
        wb.close()


def _col_index(header, name):
    if not name:
        return None
    for i, h in enumerate(header):
        if h == name:
            return i
    stripped = name.replace(" ", "")
    for i, h in enumerate(header):
        if h.replace(" ", "") == stripped:
            return i
    return None


def _to_float(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _to_str(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


# ------------------------- streaming reader (lxml) -------------------------
def _tag(name):
    return _NS + name


def _col_index_from_ref(ref):
    n = 0
    for ch in ref or "":
        if ch.isalpha():
            n = n * 26 + (ord(ch.upper()) - 64)
        else:
            break
    return n - 1


def _load_shared_strings(z):
    strings = []
    try:
        with z.open("xl/sharedStrings.xml") as f:
            for _, el in etree.iterparse(f, events=("end",), tag=_tag("si")):
                strings.append("".join(t.text or "" for t in el.iter(_tag("t"))))
                el.clear()
    except KeyError:
        pass
    return strings


def _load_date_styles(z):
    date_xf = set()
    try:
        with z.open("xl/styles.xml") as f:
            root = etree.parse(f).getroot()
        fmts = {}
        for nf in root.iter(_tag("numFmt")):
            fmts[int(nf.get("numFmtId"))] = (nf.get("formatCode") or "").lower()
        cellxfs = root.find(_tag("cellXfs"))
        if cellxfs is not None:
            for i, xf in enumerate(cellxfs.findall(_tag("xf"))):
                fid = int(xf.get("numFmtId") or 0)
                code = fmts.get(fid, "")
                if fid in _BUILTIN_DATE_FMT or "y" in code or "d" in code or "mm" in code:
                    date_xf.add(i)
    except KeyError:
        pass
    return date_xf


def _serial_to_date(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if f <= 0:
        return str(v)
    return (datetime(1899, 12, 30) + timedelta(days=f)).strftime("%Y-%m-%d")


def _sheet_xml_path(z, sheet_name):
    sheets = []
    with z.open("xl/workbook.xml") as f:
        root = etree.parse(f).getroot()
    for sh in root.iter(_tag("sheet")):
        sheets.append((sh.get("name"), sh.get(_R_NS + "id")))
    rels = {}
    with z.open("xl/_rels/workbook.xml.rels") as f:
        rroot = etree.parse(f).getroot()
    for rel in rroot:
        rels[rel.get("Id")] = rel.get("Target")
    target = None
    for name, rid in sheets:
        if sheet_name and name == sheet_name:
            target = rels.get(rid)
            break
    if target is None and sheets:
        target = rels.get(sheets[0][1])
    if target is None:
        return "xl/worksheets/sheet1.xml"
    if target.startswith("/"):
        return target.lstrip("/")
    if not target.startswith("xl/"):
        return "xl/" + target
    return target


def import_timesheet():
    with JOB_LOCK:
        return _import_timesheet()


def _import_timesheet():
    cfg = load_config()
    path = cfg["datasources"].get("timesheet")
    if not path or not os.path.exists(path):
        raise FileNotFoundError("Timesheet 文件不存在，请在“数据源设置”重新指定")
    sheet = cfg.get("timesheet_sheet") or None
    mapping = cfg.get("timesheet_mapping", {})
    missing = [label for key, label in TIMESHEET_KEYS if not mapping.get(key)]
    if missing:
        raise ValueError("Timesheet 列映射不完整：" + "、".join(missing))

    start = time.time()
    z = zipfile.ZipFile(path)
    try:
        sheet_path = _sheet_xml_path(z, sheet)
        ss = _load_shared_strings(z)
        date_xf = _load_date_styles(z)

        conn = get_conn()
        try:
            conn.execute("DELETE FROM ts_rows")
            insert = "INSERT INTO ts_rows VALUES (?,?,?,?,?,?,?)"
            cols = {}
            wanted = set()
            header_done = False
            batch = []
            total = 0
            with z.open(sheet_path) as f:
                for _, row in etree.iterparse(f, events=("end",), tag=_tag("row")):
                    if not header_done:
                        header = {}
                        for c in row.findall(_tag("c")):
                            header[_col_index_from_ref(c.get("r"))] = _cell_text(c, ss, date_xf)
                        if header:
                            width = max(header) + 1
                            for key, label in TIMESHEET_KEYS:
                                idx = _col_index([header.get(i, "") for i in range(width)], mapping.get(key))
                                if idx is None:
                                    raise ValueError("在 Timesheet 表头里找不到列：" + str(mapping.get(key)))
                                cols[key] = idx
                            wanted = set(cols.values())
                            header_done = True
                    else:
                        values = {}
                        for c in row.findall(_tag("c")):
                            ci = _col_index_from_ref(c.get("r"))
                            if ci in wanted:
                                values[ci] = _cell_text(c, ss, date_xf)
                        code = values.get(cols["project_code"])
                        if code is not None and str(code).strip() != "":
                            batch.append((
                                norm_code(code),
                                _to_str(values.get(cols["project_name"])),
                                _to_str(values.get(cols["staff"])),
                                _to_str(values.get(cols["manager"])),
                                _to_float(values.get(cols["hours"])),
                                _to_str(values.get(cols["period"])),
                                _to_str(values.get(cols["description"])),
                            ))
                            if len(batch) >= 5000:
                                conn.executemany(insert, batch)
                                total += len(batch)
                                batch = []
                    row.clear()
                    while row.getprevious() is not None:
                        del row.getparent()[0]
            if batch:
                conn.executemany(insert, batch)
                total += len(batch)
            if not header_done:
                raise ValueError("Timesheet 表头为空，请检查文件")
            _set_meta(conn, "ts_last_import", datetime.now().isoformat(timespec="seconds"))
            conn.commit()
        finally:
            conn.close()
    finally:
        z.close()
    return {"ts_rows": total, "seconds": round(time.time() - start, 1)}


def _cell_text(c, ss, date_xf):
    ttype = c.get("t")
    sattr = c.get("s")
    if ttype == "inlineStr":
        return "".join(t.text or "" for t in c.iter(_tag("t")))
    v = c.find(_tag("v"))
    txt = v.text if v is not None else None
    if ttype == "s" and txt is not None:
        try:
            return ss[int(txt)]
        except (ValueError, IndexError):
            return txt
    if txt is None:
        return ""
    if sattr is not None and date_xf:
        try:
            if int(sattr) in date_xf:
                return _serial_to_date(txt)
        except ValueError:
            pass
    return txt


def import_budget():
    with JOB_LOCK:
        return _import_budget()


def _import_budget():
    cfg = load_config()
    path = cfg["datasources"].get("wp_service")
    if not path or not os.path.exists(path):
        raise FileNotFoundError("WP服务单文件不存在，请在“数据源设置”重新指定")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = find_sheet(wb.sheetnames, "WP服务单")
        if not sheet:
            raise ValueError("WP服务单文件里找不到“WP服务单”sheet")
        ws = wb[sheet]
        header = None
        header_idx = 1
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            if row and any(v is not None and str(v).strip() for v in row):
                header = [str(v).strip() if v is not None else "" for v in row]
                header_idx = i
                break
        c_code = _col_index(header, "Engagement Code")
        c_name = _col_index(header, "Engagement Name")
        c_out = _col_index(header, "Outlook Hours")
        if c_code is None or c_out is None:
            raise ValueError("WP服务单表头缺少 Engagement Code 或 Outlook Hours")
        now = datetime.now().isoformat(timespec="seconds")
        rows = []
        for row in ws.iter_rows(min_row=header_idx + 1, values_only=True):
            if row is None or c_code >= len(row):
                continue
            code = row[c_code]
            if code is None or str(code).strip() == "":
                continue
            name = row[c_name] if c_name is not None and c_name < len(row) else ""
            out = row[c_out] if c_out < len(row) else None
            rows.append((norm_code(code), _to_str(name), _to_float(out), now))
        conn = get_conn()
        try:
            conn.execute("DELETE FROM wp_budget")
            conn.executemany("INSERT INTO wp_budget VALUES (?,?,?,?)", rows)
            _set_meta(conn, "budget_last_import", now)
            conn.commit()
        finally:
            conn.close()
    finally:
        wb.close()
    return {"budgets": len(rows)}


def _set_meta(conn, key, value):
    conn.execute(
        "INSERT INTO meta(key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, value),
    )


def get_meta(key, default=None):
    conn = get_conn()
    try:
        row = conn.execute("SELECT value FROM meta WHERE key=?", (key,)).fetchone()
        return row["value"] if row else default
    finally:
        conn.close()


def verify_datasources():
    cfg = load_config()
    out = {}
    for key, path in cfg.get("datasources", {}).items():
        out[key] = bool(path) and os.path.exists(path)
    return out
